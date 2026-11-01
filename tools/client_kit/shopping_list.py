"""Client shopping-list workbook (Hebrew, RTL) with live totals and a budget check.

Usage: shopping_list.py OUT.xlsx [ROWS]
"""
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

CREAM, ROSE, SAGE, CHOC, LINE = "F5EDE1", "E9D3CB", "D7DFCC", "4A3226", "E2D6C8"
CATS = ["רהיטים", "תאורה", "טקסטיל", "אחסון", "אקססוריז", "קירות וצבע", "אחר"]
COLS = [  # header, width
    ("#", 5), ("קטגוריה", 13), ("פריט", 26), ("חנות", 16), ("קישור", 30), ("מידות (ר×ע×ג)", 15),
    ("כמות", 7), ("מחיר ליחידה ₪", 13), ("סה״כ ₪", 12), ("חלופה חסכונית", 24),
    ("מחיר חלופה ₪", 13), ("סה״כ עם חלופה ₪", 15), ("הערות", 28),
]

font = lambda **k: Font(name="Heebo", color=CHOC, **k)
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color=LINE)
money = '#,##0 "₪"'


def build(path, rows=30):
    wb = Workbook()
    ws = wb.active
    ws.title = "רשימת קניות"
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False

    ws["A1"] = "רשימת קניות · החדר של {{שם}}"
    ws["A1"].font = Font(name="Frank Ruhl Libre", size=18, bold=True, color=CHOC)
    ws["A2"] = "Room Story Tell · המחירים נכונים ליום המסירה ועשויים להשתנות. מומלץ לבדוק מידות מול החדר לפני הזמנה."
    ws["A2"].font = font(size=9, italic=True)

    # budget box
    ws["J1"], ws["J2"], ws["J3"] = "תקציב הלקוחה", "סה״כ ברשימה", "נשאר"
    for c in ("J1", "J2", "J3"):
        ws[c].font = font(bold=True)
        ws[c].alignment = Alignment(horizontal="right")
    hdr = 5
    first, last = hdr + 1, hdr + rows
    ws["K1"] = 10000
    ws["K2"] = f"=SUM(I{first}:I{last})"
    ws["K3"] = "=K1-K2"
    for c in ("K1", "K2", "K3"):
        ws[c].number_format = money
        ws[c].font = font(bold=True)
        ws[c].fill = fill(CREAM)
    ws["L1"] = "← לעדכן"
    ws["L1"].font = font(size=9, italic=True)
    ws.conditional_formatting.add("K3", CellIsRule(operator="lessThan", formula=["0"],
                                                   fill=fill(ROSE), font=Font(color="A0665F", bold=True)))
    ws.conditional_formatting.add("K3", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=fill(SAGE)))

    for i, (h, w) in enumerate(COLS, 1):
        c = ws.cell(hdr, i, h)
        c.font = font(bold=True)
        c.fill = fill(ROSE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=Side(style="medium", color=CHOC))
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[hdr].height = 30

    cat_dv = DataValidation(type="list", formula1='"' + ",".join(CATS) + '"', allow_blank=True)
    ws.add_data_validation(cat_dv)
    for r in range(first, last + 1):
        ws.cell(r, 1, r - hdr)
        ws.cell(r, 7, None)
        ws.cell(r, 9, f'=IF(H{r}="","",H{r}*IF(G{r}="",1,G{r}))')
        ws.cell(r, 12, f'=IF(I{r}="","",IF(K{r}="",H{r},K{r})*IF(G{r}="",1,G{r}))')
        cat_dv.add(ws.cell(r, 2))
        for col in range(1, len(COLS) + 1):
            c = ws.cell(r, col)
            c.font = font(size=10)
            c.border = Border(bottom=thin)
            c.alignment = Alignment(vertical="center", wrap_text=col in (3, 10, 13))
            if col in (8, 9, 11, 12):
                c.number_format = money
        if r % 2 == 0:
            for col in range(1, len(COLS) + 1):
                ws.cell(r, col).fill = fill("FBF7F1")

    tot = last + 1
    ws.cell(tot, 3, "סה״כ").font = font(bold=True, size=12)
    for col, letter in ((9, "I"), (12, "L")):
        c = ws.cell(tot, col, f"=SUM({letter}{first}:{letter}{last})")
        c.number_format = money
        c.font = font(bold=True, size=12)
        c.fill = fill(SAGE)
    ws.cell(tot + 1, 10, "חיסכון עם החלופות").font = font(size=10)
    c = ws.cell(tot + 1, 12, f"=I{tot}-L{tot}")
    c.number_format = money
    c.font = font(size=10, bold=True)
    # highlight rows whose link cell is empty but have a price (missing link)
    ws.conditional_formatting.add(f"E{first}:E{last}", FormulaRule(formula=[f'AND($H{first}<>"",$E{first}="")'], fill=fill("F3E3B8")))
    ws.freeze_panes = ws.cell(first, 1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hdr}:{hdr}"

    # summary sheet
    s = wb.create_sheet("סיכום לפי קטגוריה")
    s.sheet_view.rightToLeft = True
    s.sheet_view.showGridLines = False
    for i, h in enumerate(("קטגוריה", "סה״כ ₪", "עם חלופות ₪"), 1):
        c = s.cell(1, i, h)
        c.font = font(bold=True)
        c.fill = fill(ROSE)
    ref = "'רשימת קניות'!"
    for j, cat in enumerate(CATS, 2):
        s.cell(j, 1, cat).font = font()
        for col, letter in ((2, "I"), (3, "L")):
            c = s.cell(j, col, f'=SUMIF({ref}$B${first}:$B${last},A{j},{ref}{letter}${first}:{letter}${last})')
            c.number_format = money
            c.font = font()
    end = len(CATS) + 2
    s.cell(end, 1, "סה״כ").font = font(bold=True)
    for col, letter in ((2, "B"), (3, "C")):
        c = s.cell(end, col, f"=SUM({letter}2:{letter}{end - 1})")
        c.number_format = money
        c.font = font(bold=True)
        c.fill = fill(SAGE)
    for letter, w in (("A", 16), ("B", 14), ("C", 16)):
        s.column_dimensions[letter].width = w

    wb.save(path)


if __name__ == "__main__":
    build(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else 30)
    print("wrote", sys.argv[1])
